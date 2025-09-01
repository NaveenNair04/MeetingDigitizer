import json
import time
import base64
import os
from datetime import datetime
from kafka import KafkaConsumer
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import OxmlElement, qn
import threading
from queue import PriorityQueue
import logging

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class MultiTopicKafkaConsumer:
    def __init__(self, bootstrap_servers=['kafka:9092'], output_dir='./output'):
        self.bootstrap_servers = bootstrap_servers
        self.output_dir = output_dir
        self.diagram_dir = os.path.join(output_dir, 'diagrams')
        
        # Create output directories
        os.makedirs(self.output_dir, exist_ok=True)
        os.makedirs(self.diagram_dir, exist_ok=True)
        
        # Store all messages in memory for sorting
        self.all_messages = []
        
        # Fixed document filename
        self.doc_filename = os.path.join(self.output_dir, 'kafka_consumer_data.docx')
        
        # Topics to consume from
        self.topics = ['audio-transcripts', 'ocr-sentences', 'diagram-detections']
        
        # Consumer configuration
        self.consumer = KafkaConsumer(
            *self.topics,
            bootstrap_servers=self.bootstrap_servers,
            auto_offset_reset='latest',  # Change to 'earliest' to consume from beginning
            enable_auto_commit=True,
            group_id='multi-topic-consumer-group',
            value_deserializer=lambda x: json.loads(x.decode('utf-8'))
        )
        
        # Track processed messages count
        self.processed_count = 0
        
    def get_timestamp_from_message(self, message_data, topic):
        """Extract timestamp from message based on topic"""
        try:
            if topic == 'audio-transcripts':
                return message_data.get('kafka_timestamp', time.time())
            elif topic == 'ocr-sentences':
                return message_data.get('detection_timestamp', time.time())
            elif topic == 'diagram-detections':
                return message_data.get('detection_timestamp', time.time())
            else:
                return time.time()
        except:
            return time.time()
    
    def save_diagram_image(self, diagram_b64, frame_id, diagram_type):
        """Save diagram image to file"""
        try:
            if diagram_b64:
                # Decode base64 image
                image_data = base64.b64decode(diagram_b64)
                
                # Create filename
                timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                filename = f"diagram_{frame_id}_{diagram_type}_{timestamp}.png"
                filepath = os.path.join(self.diagram_dir, filename)
                
                # Save image
                with open(filepath, 'wb') as f:
                    f.write(image_data)
                
                logger.info(f"Saved diagram image: {filename}")
                return filepath
        except Exception as e:
            logger.error(f"Error saving diagram image: {e}")
            return None
    
    def create_document_from_all_messages(self):
        """Create a new Word document with all messages sorted by timestamp"""
        # Create new document
        doc = Document()
        doc.add_heading('Multi-Source Data Stream', 0)
        
        # Add document info
        p = doc.add_paragraph()
        p.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").bold = True
        p = doc.add_paragraph(f"Total Messages: {len(self.all_messages)}")
        doc.add_page_break()
        
        # Sort all messages by timestamp
        sorted_messages = sorted(self.all_messages, key=lambda x: x[0])
        
        # Process each message in timestamp order
        for timestamp, topic, message_data in sorted_messages:
            if topic == 'audio-transcripts':
                self.format_audio_transcript_to_doc(doc, message_data)
            elif topic == 'ocr-sentences':
                self.format_ocr_sentence_to_doc(doc, message_data)
            elif topic == 'diagram-detections':
                self.format_diagram_detection_to_doc(doc, message_data)
        
        return doc
    
    def add_separator_to_doc(self, doc):
        """Add a separator line to the document"""
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("─" * 80)
        run.font.color.rgb = None  # Default color
    
    def format_audio_transcript_to_doc(self, doc, data):
        """Format audio transcript message for Word document"""
        doc.add_heading(f'🎤 Audio Transcript', level=2)
        
        # Add timestamp
        timestamp = data.get('kafka_timestamp', time.time())
        readable_time = datetime.fromtimestamp(timestamp).strftime('%Y-%m-%d %H:%M:%S')
        p = doc.add_paragraph()
        p.add_run(f"Timestamp: {readable_time}").bold = True
        
        # Add time range
        start_time = data.get('start', 'N/A')
        end_time = data.get('end', 'N/A')
        p = doc.add_paragraph(f"Time Range: {start_time}s - {end_time}s")
        
        # Add speaker
        speaker = data.get('speaker', 'Unknown')
        p = doc.add_paragraph(f"Speaker: {speaker}")
        
        # Add transcript text
        text = data.get('text', '')
        p = doc.add_paragraph()
        p.add_run("Transcript: ").bold = True
        p.add_run(text)
        
        self.add_separator_to_doc(doc)
    
    def format_ocr_sentence_to_doc(self, doc, data):
        """Format OCR sentence message for Word document"""
        doc.add_heading(f'📄 OCR Text Detection', level=2)
        
        # Add timestamp
        timestamp = data.get('detection_timestamp', time.time())
        readable_time = datetime.fromtimestamp(timestamp).strftime('%Y-%m-%d %H:%M:%S')
        p = doc.add_paragraph()
        p.add_run(f"Timestamp: {readable_time}").bold = True
        
        # Add frame info
        frame_id = data.get('frame_id', 'N/A')
        p = doc.add_paragraph(f"Frame ID: {frame_id}")
        
        # Add line number
        line_number = data.get('line_number', 'N/A')
        p = doc.add_paragraph(f"Line Number: {line_number}")
        
        # Add confidence and word count
        confidence = data.get('confidence', 0)
        word_count = data.get('word_count', 0)
        p = doc.add_paragraph(f"Confidence: {confidence:.2f} | Word Count: {word_count}")
        
        # Add detected text
        sentence = data.get('sentence', '')
        p = doc.add_paragraph()
        p.add_run("Detected Text: ").bold = True
        p.add_run(sentence)
        
        self.add_separator_to_doc(doc)
    
    def format_diagram_detection_to_doc(self, doc, data):
        """Format diagram detection message for Word document"""
        doc.add_heading(f'📊 Diagram Detection', level=2)
        
        # Add timestamp
        timestamp = data.get('detection_timestamp', time.time())
        readable_time = datetime.fromtimestamp(timestamp).strftime('%Y-%m-%d %H:%M:%S')
        p = doc.add_paragraph()
        p.add_run(f"Timestamp: {readable_time}").bold = True
        
        # Add frame info
        frame_id = data.get('frame_id', 'N/A')
        diagram_type = data.get('diagram_type', 'Unknown')
        p = doc.add_paragraph(f"Frame ID: {frame_id} | Type: {diagram_type}")
        
        # Add bounding box and dimensions
        bbox = data.get('bbox', [])
        area = data.get('area', 0)
        aspect_ratio = data.get('aspect_ratio', 0)
        p = doc.add_paragraph(f"Bounding Box: {bbox}")
        p = doc.add_paragraph(f"Area: {area} | Aspect Ratio: {aspect_ratio:.2f}")
        
        # Add confidence
        confidence = data.get('confidence', 0)
        p = doc.add_paragraph(f"Confidence: {confidence:.2f}")
        
        # Add extracted text
        extracted_text = data.get('extracted_text', [])
        if extracted_text:
            p = doc.add_paragraph()
            p.add_run("Extracted Text: ").bold = True
            p.add_run(str(extracted_text))
        
        # Save and reference diagram image
        diagram_b64 = data.get('diagram_image')
        if diagram_b64:
            image_path = self.save_diagram_image(diagram_b64, frame_id, diagram_type)
            if image_path:
                p = doc.add_paragraph(f"Diagram saved to: {os.path.basename(image_path)}")
                
                # Try to add image to document (optional, in case of size issues)
                try:
                    # Add image with reasonable size
                    doc.add_picture(image_path, width=Inches(4))
                except Exception as e:
                    logger.warning(f"Could not embed image in document: {e}")
                    p = doc.add_paragraph(f"Image file: {os.path.basename(image_path)}")
        
        self.add_separator_to_doc(doc)
    
    def rebuild_and_save_document(self):
        """Rebuild the entire document with all messages and save (overwrite)"""
        try:
            # Create new document
            doc = self.create_document_from_all_messages()
            
            # Save document (overwrite existing)
            doc.save(self.doc_filename)
            logger.info(f"Document updated with {len(self.all_messages)} messages: {self.doc_filename}")
            return self.doc_filename
            
        except Exception as e:
            logger.error(f"Error saving document: {e}")
            return None
    
    def consume_messages(self, save_interval=30):
        """
        Main consumer loop - saves to single document file, overwriting each time
        
        Args:
            save_interval: Interval in seconds to rebuild and save the document
        """
        logger.info(f"Starting consumer for topics: {self.topics}")
        logger.info(f"Output file: {self.doc_filename}")
        logger.info("Document will be overwritten with updated data every save interval")
        
        last_save_time = time.time()
        
        try:
            for message in self.consumer:
                try:
                    topic = message.topic
                    message_data = message.value
                    
                    # Get timestamp for ordering
                    timestamp = self.get_timestamp_from_message(message_data, topic)
                    
                    # Add message to our collection
                    self.all_messages.append((timestamp, topic, message_data))
                    self.processed_count += 1
                    
                    logger.info(f"Received message #{self.processed_count} from {topic}")
                    
                    # Save document periodically (overwrite)
                    current_time = time.time()
                    if current_time - last_save_time >= save_interval:
                        self.rebuild_and_save_document()
                        last_save_time = current_time
                
                except Exception as e:
                    logger.error(f"Error processing message: {e}")
                    continue
        
        except KeyboardInterrupt:
            logger.info("Consumer interrupted by user")
        
        finally:
            # Final save with all collected messages
            final_path = self.rebuild_and_save_document()
            logger.info(f"Final document saved to: {final_path}")
            logger.info(f"Total messages processed: {self.processed_count}")
            
            # Close consumer
            self.consumer.close()
            logger.info("Consumer closed")

def main():
    """Main function to run the consumer"""
    # Configuration
    KAFKA_SERVERS = ['kafka:9092']  # Update with your Kafka servers
    OUTPUT_DIR = './output'       # Update with your desired output directory
    
    # Create consumer instance
    consumer = MultiTopicKafkaConsumer(
        bootstrap_servers=KAFKA_SERVERS,
        output_dir=OUTPUT_DIR
    )
    
    # Start consuming
    logger.info("Starting Kafka consumer...")
    consumer.consume_messages(
        save_interval=30    # Rebuild and save document every 30 seconds
    )

if __name__ == "__main__":
    main()