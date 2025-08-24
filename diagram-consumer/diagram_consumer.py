import json
import base64
import cv2
import numpy as np
import os
from kafka import KafkaConsumer
from sklearn.cluster import DBSCAN
from docx import Document
from docx.shared import Inches

TOPIC = "video-stream"
BOOTSTRAP_SERVERS = "kafka:9092"
GROUP_ID = "diagram-consumer"

OUTPUT_FOLDER = "output_diagram"
CROPPED_FOLDER = os.path.join(OUTPUT_FOLDER, "cropped")
os.makedirs(CROPPED_FOLDER, exist_ok=True)

doc = Document()
doc.add_heading("Detected Diagrams", level=1)
previous_diagrams = []

def detect_diagrams(image, min_area=5000, margin=60):
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    blurred = cv2.GaussianBlur(gray, (5,5), 0)
    edges = cv2.Canny(blurred, 50, 150)
    contours, _ = cv2.findContours(edges, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    boxes = []
    for cnt in contours:
        epsilon = 0.02*cv2.arcLength(cnt, True)
        approx = cv2.approxPolyDP(cnt, epsilon, True)
        x, y, w, h = cv2.boundingRect(cnt)
        if w*h >= min_area and len(approx) >=4:
            x_ext = max(x-margin,0)
            y_ext = max(y-margin,0)
            x2_ext = min(x+w+margin, image.shape[1])
            y2_ext = min(y+h+margin, image.shape[0])
            boxes.append((x_ext,y_ext,x2_ext,y2_ext))
    if not boxes:
        return []

    centers = np.array([[(x1+x2)/2, (y1+y2)/2] for x1,y1,x2,y2 in boxes])
    clustering = DBSCAN(eps=100, min_samples=1).fit(centers)
    labels = clustering.labels_

    final_boxes = []
    for label in set(labels):
        cluster_boxes = np.array([b for b,l in zip(boxes, labels) if l==label])
        x1 = int(np.min(cluster_boxes[:,0]))
        y1 = int(np.min(cluster_boxes[:,1]))
        x2 = int(np.max(cluster_boxes[:,2]))
        y2 = int(np.max(cluster_boxes[:,3]))
        final_boxes.append((x1,y1,x2,y2))
    return final_boxes

def iou(box1, box2):
    x1,y1,x2,y2 = box1
    x3,y3,x4,y4 = box2
    xi1, yi1 = max(x1,x3), max(y1,y3)
    xi2, yi2 = min(x2,x4), min(y2,y4)
    inter_area = max(0, xi2-xi1) * max(0, yi2-yi1)
    union_area = (x2-x1)*(y2-y1) + (x4-x3)*(y4-y3) - inter_area
    return inter_area / union_area if union_area>0 else 0

consumer = KafkaConsumer(
    TOPIC,
    bootstrap_servers=BOOTSTRAP_SERVERS,
    group_id=GROUP_ID,
    auto_offset_reset="latest",
    value_deserializer=lambda m: json.loads(m.decode("utf-8")),
)

frame_count = 0
print("🟢 Waiting for frames for Diagram detection...")

for msg in consumer:
    try:
        payload = msg.value
        timestamp = payload.get("timestamp")
        chunk = base64.b64decode(payload["data"])

        nparr = np.frombuffer(chunk, np.uint8)
        frame = cv2.imdecode(nparr, cv2.IMREAD_COLOR)
        if frame is None:
            continue

        diagrams = detect_diagrams(frame)
        for i, box in enumerate(diagrams, start=1):
            if all(iou(box, prev) < 0.7 for prev in previous_diagrams):
                previous_diagrams.append(box)
                x1,y1,x2,y2 = box
                crop = frame[y1:y2, x1:x2]
                crop_file = os.path.join(CROPPED_FOLDER, f"frame{frame_count}_diagram{i}.jpg")
                cv2.imwrite(crop_file, crop)
                doc.add_paragraph(f"Frame {frame_count} (t={timestamp:.2f}s) - Diagram {i}")
                doc.add_picture(crop_file, width=Inches(4))

        frame_count += 1
        if frame_count % 5 == 0:
            doc.save(os.path.join(OUTPUT_FOLDER, "Detected_Diagrams.docx"))
            print(f"💾 Saved diagrams after {frame_count} frames")

    except Exception as e:
        print(f"Error processing frame: {e}")
